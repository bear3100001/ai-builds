#!/usr/bin/env python3
"""Inspect requirements/index workbooks and build a validated finishes schedule."""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REQ_SHEET = "Finish Requirements"
OUT_SHEET = "Finishes Schedule"
REQ_REQUIRED = ("Location", "Description")
REQ_OPTIONAL = ("Colour Reference", "Finish", "Notes")
INDEX_COLUMNS = (
    "Code", "Brief Description", "Dimensions", "Location",
    "Material Description", "Image", "Colour Reference", "Finish",
    "Acoustic Requirements", "Confidence", "Check notes",
)
OUTPUT_COLUMNS = (
    "Code", "Brief Description", "Dimensions", "Location",
    "Material Description", "Image", "Colour Reference", "Finish",
    "Acoustic Requirements", "Alternatives", "Check notes",
)
CATEGORY_ORDER = {"FL": 0, "LIN": 1, "PLS": 2, "PT": 3, "TRM": 4}
CONFIDENCE_ORDER = {"high": 0, "medium": 1, "low": 2}
PINK = "F4CCCC"
HEADER = "D9EAF7"
CATEGORY_WORDS = {
    "FL": {"floor", "flooring", "carpet", "tile", "stone", "marble", "terrazzo", "vinyl", "resin"},
    "LIN": {"wall", "lining", "wallcovering", "wallpaper", "cladding", "panel", "tile"},
    "PLS": {"plasterboard", "drywall", "gypsum", "board"},
    "PT": {"paint", "coating", "emulsion"},
    "TRM": {"skirting", "baseboard", "trim", "profile", "edging"},
}
STOPWORDS = {
    "a", "an", "and", "as", "at", "be", "by", "for", "from", "in", "is", "it", "of", "on",
    "or", "the", "to", "with", "finish", "finishes", "required", "provide", "supply", "project",
}


class InputProblem(ValueError):
    pass


def text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def norm(value) -> str:
    value = text(value).casefold().replace("matte", "matt")
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return " ".join(value.split())


def tokens(value) -> set[str]:
    return {w for w in norm(value).split() if len(w) > 1 and w not in STOPWORDS}


def partial_match(requested, indexed) -> bool:
    req = norm(requested)
    idx = norm(indexed)
    return not req or bool(idx and (req in idx or idx in req))


def header_map(ws, required: tuple[str, ...], allowed: tuple[str, ...] | None = None):
    for row_no in range(1, min(ws.max_row, 30) + 1):
        values = [text(ws.cell(row_no, c).value) for c in range(1, ws.max_column + 1)]
        mapping = {v: i + 1 for i, v in enumerate(values) if v}
        if all(name in mapping for name in required):
            if allowed is not None:
                unknown = [v for v in values if v and v not in allowed]
                if unknown:
                    raise InputProblem("Wrong headers. Unexpected column(s): " + ", ".join(unknown))
            return row_no, mapping
    raise InputProblem("Wrong headers. Required column(s): " + ", ".join(required))


def load_requirements(path: Path):
    try:
        wb = load_workbook(path, data_only=False)
    except Exception as exc:
        raise InputProblem(f"The Finish Requirements workbook could not be opened: {exc}") from exc
    if wb.sheetnames != [REQ_SHEET]:
        raise InputProblem(f'The template must contain exactly one worksheet named "{REQ_SHEET}".')
    ws = wb[REQ_SHEET]
    header_row, cols = header_map(ws, REQ_REQUIRED, REQ_REQUIRED + REQ_OPTIONAL)
    rows = []
    for row_no in range(header_row + 1, ws.max_row + 1):
        record = {
            name: text(ws.cell(row_no, cols[name]).value) if name in cols else ""
            for name in REQ_REQUIRED + REQ_OPTIONAL
        }
        if not any(record.values()):
            continue
        marker = norm(" ".join(record.values()))
        if ("example" in marker and ("delete" in marker or "remove" in marker)) or "delete this row" in marker:
            continue
        rows.append({"source_row": row_no, **record})
    if not rows:
        raise InputProblem("The Finish Requirements worksheet contains no requirement rows.")
    return rows


def load_index(path: Path):
    try:
        wb = load_workbook(path, data_only=False)
    except Exception as exc:
        raise InputProblem(f"The Product Index could not be opened: {exc}") from exc
    for ws in wb.worksheets:
        try:
            header_row, cols = header_map(ws, ("Code", "Brief Description", "Location"))
            break
        except InputProblem:
            continue
    else:
        raise InputProblem("The Product Index does not contain the required headers.")
    missing = [c for c in INDEX_COLUMNS if c not in cols]
    if missing:
        raise InputProblem("The Product Index is missing column(s): " + ", ".join(missing))
    rows = []
    seen = set()
    for row_no in range(header_row + 1, ws.max_row + 1):
        record = {name: text(ws.cell(row_no, cols[name]).value) for name in INDEX_COLUMNS}
        code = record["Code"]
        if not code:
            continue
        if not re.fullmatch(r"(FL|LIN|PLS|PT|TRM)-\d+-[A-Za-z0-9]+", code, re.I):
            continue
        key = code.casefold()
        if key in seen:
            raise InputProblem(f"The Product Index contains duplicate code {code}.")
        seen.add(key)
        record["_row"] = row_no
        record["_category"] = code.split("-", 1)[0].upper()
        rows.append(record)
    if not rows:
        raise InputProblem("The Product Index contains no usable finish products.")
    return rows


def infer_categories(description: str) -> list[str]:
    words = tokens(description)
    ranked = []
    for category, category_words in CATEGORY_WORDS.items():
        score = len(words & category_words)
        if score:
            ranked.append((score, category))
    ranked.sort(key=lambda item: (-item[0], CATEGORY_ORDER[item[1]]))
    return [category for _, category in ranked]


def semantic_score(req, product) -> float:
    requested = tokens(" ".join((req["Description"], req["Notes"])))
    product_words = tokens(" ".join(text(product.get(c)) for c in INDEX_COLUMNS))
    if not requested:
        return 0.0
    overlap = len(requested & product_words)
    return overlap / max(len(requested), 1)


def location_score(location: str, product_location: str) -> int:
    req = tokens(location)
    idx = tokens(product_location)
    if not req or not idx:
        return 2
    if req & idx:
        return 0
    nreq, nidx = norm(location), norm(product_location)
    return 0 if nreq in nidx or nidx in nreq else 1


def candidates_for(req, products):
    categories = infer_categories(req["Description"] + " " + req["Notes"])
    eligible = []
    for product in products:
        if req["Colour Reference"] and not partial_match(req["Colour Reference"], product["Colour Reference"]):
            continue
        if req["Finish"] and not partial_match(req["Finish"], product["Finish"]):
            continue
        category_penalty = 0 if not categories or product["_category"] in categories else 1
        loc = location_score(req["Location"], product["Location"])
        conf = CONFIDENCE_ORDER.get(norm(product["Confidence"]), 3)
        sem = semantic_score(req, product)
        eligible.append((category_penalty, loc, conf, -sem, product["Code"], product))
    eligible.sort(key=lambda item: item[:5])
    return [item[-1] for item in eligible[:50]]


def public_product(product):
    return {name: product[name] for name in INDEX_COLUMNS}


def inspect_command(args):
    requirements = load_requirements(Path(args.requirements))
    products = load_index(Path(args.index))
    payload = {
        "requirements": [],
        "index_product_count": len(products),
        "rules": {
            "colour_and_finish_are_mandatory_when_requested": True,
            "blank_index_values_never_match_populated_filters": True,
            "maximum_alternatives": 2,
        },
    }
    for req in requirements:
        status_hint = "INPUT ERROR" if not req["Location"] or not req["Description"] else "REVIEW"
        payload["requirements"].append({
            **req,
            "status_hint": status_hint,
            "inferred_categories": infer_categories(req["Description"] + " " + req["Notes"]),
            "candidates": [public_product(p) for p in candidates_for(req, products)] if status_hint == "REVIEW" else [],
        })
    Path(args.output).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def load_decisions(path: Path):
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        raise InputProblem(f"The decisions JSON could not be read: {exc}") from exc
    rows = payload.get("rows")
    if not isinstance(rows, list):
        raise InputProblem('The decisions JSON must contain a "rows" list.')
    return rows


def family_of(index_code: str) -> tuple[str, str]:
    parts = index_code.split("-")
    return parts[0].upper(), parts[1]


def excel_letter(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def validate_and_prepare(requirements, products, decisions):
    product_by_code = {p["Code"].casefold(): p for p in products}
    requirement_by_row = {r["source_row"]: r for r in requirements}
    decision_by_row = {}
    allowed_statuses = {"MATCH", "NO MATCH", "INPUT ERROR"}
    for decision in decisions:
        source_row = decision.get("source_row")
        if source_row not in requirement_by_row:
            raise InputProblem(f"Decision refers to unknown source_row {source_row}.")
        if source_row in decision_by_row:
            raise InputProblem(f"Duplicate decision for source_row {source_row}.")
        status = text(decision.get("status")).upper()
        if status not in allowed_statuses:
            raise InputProblem(f"Invalid status for source_row {source_row}: {status}")
        decision_by_row[source_row] = {**decision, "status": status}
    missing = sorted(set(requirement_by_row) - set(decision_by_row))
    if missing:
        raise InputProblem("Missing decision(s) for source_row: " + ", ".join(map(str, missing)))

    prepared = []
    for req in requirements:
        decision = decision_by_row[req["source_row"]]
        status = decision["status"]
        missing_required = [name for name in REQ_REQUIRED if not req[name]]
        if missing_required and status != "INPUT ERROR":
            raise InputProblem(
                f"source_row {req['source_row']} must be INPUT ERROR because {', '.join(missing_required)} is blank."
            )
        if not missing_required and status == "INPUT ERROR" and not text(decision.get("reason")):
            raise InputProblem(f"source_row {req['source_row']} needs an INPUT ERROR reason.")
        selected = None
        alternatives = []
        if status == "MATCH":
            code = text(decision.get("selected_index_code"))
            selected = product_by_code.get(code.casefold())
            if selected is None:
                raise InputProblem(
                    f"source_row {req['source_row']} selects a code not in the current Product Index: {code}"
                )
            if req["Colour Reference"] and not partial_match(req["Colour Reference"], selected["Colour Reference"]):
                raise InputProblem(f"source_row {req['source_row']} selected product fails the Colour Reference filter.")
            if req["Finish"] and not partial_match(req["Finish"], selected["Finish"]):
                raise InputProblem(f"source_row {req['source_row']} selected product fails the Finish filter.")
            alt_codes = decision.get("alternative_index_codes") or []
            if not isinstance(alt_codes, list) or len(alt_codes) > 2:
                raise InputProblem(f"source_row {req['source_row']} may have at most two alternatives.")
            seen = {selected["Code"].casefold()}
            for alt_code in alt_codes:
                alt = product_by_code.get(text(alt_code).casefold())
                if alt is None:
                    raise InputProblem(
                        f"source_row {req['source_row']} alternative is not in the current Product Index: {alt_code}"
                    )
                if alt["Code"].casefold() in seen:
                    raise InputProblem(f"source_row {req['source_row']} repeats a selected or alternative code.")
                if alt["_category"] != selected["_category"]:
                    raise InputProblem(
                        f"source_row {req['source_row']} alternative {alt['Code']} is from a different category."
                    )
                if req["Colour Reference"] and not partial_match(req["Colour Reference"], alt["Colour Reference"]):
                    raise InputProblem(
                        f"source_row {req['source_row']} alternative {alt['Code']} fails the Colour Reference filter."
                    )
                if req["Finish"] and not partial_match(req["Finish"], alt["Finish"]):
                    raise InputProblem(
                        f"source_row {req['source_row']} alternative {alt['Code']} fails the Finish filter."
                    )
                seen.add(alt["Code"].casefold())
                alternatives.append(alt)
        elif text(decision.get("selected_index_code")):
            raise InputProblem(f"source_row {req['source_row']} cannot select a product with status {status}.")
        prepared.append({"req": req, "decision": decision, "selected": selected, "alternatives": alternatives})
    return prepared


def assign_project_codes(prepared):
    next_family = defaultdict(lambda: 101)
    family_numbers = {}
    variant_numbers = defaultdict(dict)
    project_codes = {}
    for item in prepared:
        product = item["selected"]
        if not product:
            continue
        category, family = family_of(product["Code"])
        family_key = (category, family)
        if family_key not in family_numbers:
            family_numbers[family_key] = next_family[category]
            next_family[category] += 1
        code_key = product["Code"].casefold()
        if code_key not in variant_numbers[family_key]:
            variant_numbers[family_key][code_key] = len(variant_numbers[family_key]) + 1
        project_codes[code_key] = (
            f"{category}-{family_numbers[family_key]:03d}-{excel_letter(variant_numbers[family_key][code_key])}"
        )
    return project_codes


def build_command(args):
    requirements = load_requirements(Path(args.requirements))
    products = load_index(Path(args.index))
    decisions = load_decisions(Path(args.decisions))
    prepared = validate_and_prepare(requirements, products, decisions)
    project_codes = assign_project_codes(prepared)

    output_rows = []
    for item in prepared:
        req, decision, product = item["req"], item["decision"], item["selected"]
        if product:
            project_code = project_codes[product["Code"].casefold()]
            alt_labels = []
            for alt in item["alternatives"]:
                alt_labels.append(
                    project_codes.get(alt["Code"].casefold(), f"{alt['Code']} - {alt['Brief Description']}")
                )
            prefix = f"Index: {product['Code']}. Requested: {req['Description']}."
            extra = text(decision.get("reason"))
            notes = prefix + ((" " + extra) if extra else "")
            row = {
                "Code": project_code,
                "Brief Description": product["Brief Description"],
                "Dimensions": product["Dimensions"],
                "Location": req["Location"],
                "Material Description": product["Material Description"],
                "Image": product["Image"],
                "Colour Reference": product["Colour Reference"],
                "Finish": product["Finish"],
                "Acoustic Requirements": product["Acoustic Requirements"],
                "Alternatives": "; ".join(alt_labels),
                "Check notes": notes,
                "_sort": (CATEGORY_ORDER[product["_category"]], req["source_row"]),
                "_flag": False,
            }
        else:
            status = decision["status"]
            reason = text(decision.get("reason")) or (
                "Missing or unrecognisable input." if status == "INPUT ERROR" else "No indexed product meets the request."
            )
            row = {name: "" for name in OUTPUT_COLUMNS}
            row.update({
                "Code": status,
                "Brief Description": req["Description"],
                "Location": req["Location"],
                "Colour Reference": req["Colour Reference"],
                "Finish": req["Finish"],
                "Check notes": f"Requested: {req['Description']}. {reason}".strip(),
                "_sort": ((5 if status == "NO MATCH" else 6), req["source_row"]),
                "_flag": True,
            })
        output_rows.append(row)
    output_rows.sort(key=lambda row: row["_sort"])

    wb = Workbook()
    ws = wb.active
    ws.title = OUT_SHEET
    ws.append(OUTPUT_COLUMNS)
    for row in output_rows:
        ws.append([row[name] for name in OUTPUT_COLUMNS])
        if row["_flag"]:
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor=PINK)

    thin = Side(style="thin", color="B7C9D6")
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=HEADER)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [16, 28, 18, 24, 56, 22, 22, 18, 24, 36, 62]
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 32
    for row_no in range(2, ws.max_row + 1):
        ws.row_dimensions[row_no].height = 54

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    check = load_workbook(output, data_only=False)
    if check.sheetnames != [OUT_SHEET] or check[OUT_SHEET].max_row != len(output_rows) + 1:
        raise InputProblem("Output validation failed.")


def main():
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)

    inspect_parser = subparsers.add_parser("inspect")
    inspect_parser.add_argument("--requirements", required=True)
    inspect_parser.add_argument("--index", required=True)
    inspect_parser.add_argument("--output", required=True)
    inspect_parser.set_defaults(func=inspect_command)

    build_parser = subparsers.add_parser("build")
    build_parser.add_argument("--requirements", required=True)
    build_parser.add_argument("--index", required=True)
    build_parser.add_argument("--decisions", required=True)
    build_parser.add_argument("--output", required=True)
    build_parser.set_defaults(func=build_command)

    args = parser.parse_args()
    try:
        args.func(args)
    except InputProblem as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2)


if __name__ == "__main__":
    main()
